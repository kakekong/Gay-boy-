"""A supplier PO carries notes for the vendor — delivery, packing, terms.

They can be typed when the PO is raised or later, and save on the spot (like
the ETA): they change nothing priced or counted, so they don't wait on the
director. Each change is audited. Finance can read them but not write them.
They print on the PO the supplier receives (PDF and spreadsheet) and show on
the supplier's portal.
"""
import asyncio, io, os, sys, uuid
os.environ.update(DATABASE_URL="postgresql+asyncpg://postgres@127.0.0.1:55432/transmisi_test",
    APP_ENV="dev", DEMO_SEED_PASSWORD="test-pass-123",
    STORAGE_LOCAL_DIR="/tmp/storage_test", JWT_SECRET="e2e-test-secret")
sys.path.insert(0, "/home/user/Gay-boy-/backend")
import httpx, logging; logging.disable(logging.INFO)
TAG = uuid.uuid4().hex[:5].upper()
PASS, FAIL = [], []
def check(n, c, d=""):
    (PASS if c else FAIL).append(n); print(("  PASS " if c else "  FAIL ")+n+(f"  [{d}]" if d and not c else ""))
def J(r):
    try: return r.json()
    except Exception: return {"_": r.text[:200]}


async def main():
    from app.scripts.seed import ensure_schema; await ensure_schema()
    from app.main import app
    from app.core.db import SessionLocal
    from app.core.security import hash_password
    from app.models.audit import AuditLog
    from app.models.user import User
    from sqlalchemy import select
    c = httpx.AsyncClient(transport=httpx.ASGITransport(app=app),
                          base_url="http://t/api/v1", timeout=180)
    async def login(e):
        r = await c.post("/auth/login", json={"email": e, "password": "test-pass-123"})
        return {"Authorization": f"Bearer {r.json()['access_token']}"}
    d = await login("director@demo.local")
    pur = await login("purchasing@demo.local")
    fin = await login("finance@demo.local")
    sup = J(await c.post("/purchasing/suppliers", headers=pur, json={"name": f"PT Catatan {TAG}"}))["id"]
    async with SessionLocal() as s:
        s.add(User(email=f"sup-notes-{TAG}@portal.test".lower(), full_name="Portal supplier", role="supplier",
                   password_hash=hash_password("test-pass-123"), is_active=True,
                   linked_supplier_id=uuid.UUID(sup)))
        await s.commit()
    sp = await login(f"sup-notes-{TAG}@portal.test".lower())

    line = {"description": f"V-BELT {TAG}", "qty": 4, "uom": "pcs", "unit_price": 100_000, "amount": 400_000}
    first = f"Deliver to Gate 2 before 15:00 — {TAG}"

    print("\n── raising a PO with notes ──")
    po = J(await c.post("/purchasing/po", headers=pur, json={
        "supplier_id": sup, "po_date": "2026-10-12", "items": [line], "total": 400_000,
        "notes": f"  {first}\n"}))
    check("purchasing raises a PO with notes", bool(po.get("id")), str(po)[:200])
    got = J(await c.get(f"/purchasing/po/{po['id']}", headers=d))
    check("...the notes are kept, trimmed", got.get("notes") == first, repr(got.get("notes")))

    print("\n── editing them ──")
    second = f"Pack each belt separately & label <A-{TAG}>.\nCall Budi on arrival."
    def queued(rows):
        return [x["id"] for x in rows if x["target_type"] == "supplier_po" and x["target_id"] == po["id"]]
    waiting = queued(J(await c.get("/approvals", headers=d)))   # the new PO itself may be waiting
    r = await c.patch(f"/purchasing/po/{po['id']}", headers=pur, json={"notes": second})
    got = J(await c.get(f"/purchasing/po/{po['id']}", headers=d))
    check("purchasing's note change applies straight away", r.status_code == 200
          and got.get("notes") == second, f"{r.status_code} {J(r)} {got.get('notes')!r}")
    now = queued(J(await c.get("/approvals", headers=d)))
    check("...and nothing new is queued for the director", now == waiting, f"{waiting} -> {now}")
    async with SessionLocal() as s:
        audits = (await s.scalars(select(AuditLog).where(
            AuditLog.entity == "supplier_po", AuditLog.entity_id == uuid.UUID(po["id"])))).all()
    check("...and the change is in the audit log, old and new",
          any(a.before.get("notes") == first and a.after.get("notes") == second for a in audits),
          str([(a.before, a.after) for a in audits])[:200])

    r = await c.patch(f"/purchasing/po/{po['id']}", headers=fin, json={"notes": "finance was here"})
    got = J(await c.get(f"/purchasing/po/{po['id']}", headers=fin))
    check("finance can't write notes", r.status_code == 403, f"{r.status_code} {J(r)}")
    check("...but can read them", got.get("notes") == second, repr(got.get("notes")))

    print("\n── where the supplier sees them ──")
    r = await c.get(f"/purchasing/po/{po['id']}/export.pdf", headers=d)
    text = ""
    if r.status_code == 200:
        from pypdf import PdfReader
        text = "".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages)
    check("the PDF prints them — '&' and '<' included", r.status_code == 200
          and f"<A-{TAG}>" in text and "Call Budi" in text, f"{r.status_code} {text[-300:]!r}")
    r = await c.get(f"/purchasing/po/{po['id']}/export.xlsx", headers=d)
    cells = []
    if r.status_code == 200:
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(r.content)).active
        cells = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v]
    check("the spreadsheet carries them", r.status_code == 200 and second in cells,
          f"{r.status_code} {cells[-5:]}")
    orders = J(await c.get("/portal/supplier/orders", headers=sp))
    mine = next((p for p in orders.get("purchase_orders", []) if p["id"] == po["id"]), {})
    check("the supplier's portal shows them", mine.get("notes") == second, repr(mine.get("notes")))

    print("\n── clearing ──")
    r = await c.patch(f"/purchasing/po/{po['id']}", headers=d, json={"notes": "   "})
    got = J(await c.get(f"/purchasing/po/{po['id']}", headers=d))
    check("blank notes clear them", r.status_code == 200 and got.get("notes") is None, repr(got.get("notes")))
    r = await c.get(f"/purchasing/po/{po['id']}/export.pdf", headers=d)
    check("...and the PDF still renders without them", r.status_code == 200, str(r.status_code))

    await c.aclose()
    print(f"\n{len(PASS)} passed, {len(FAIL)} failed")
    if FAIL:
        print("FAILED: " + "; ".join(FAIL)); sys.exit(1)

asyncio.run(main())
